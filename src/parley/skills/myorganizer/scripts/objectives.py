"""Read the user's objectives spreadsheet and print it as JSON.

Part of the /myorganizer skill. Where the workbook lives and how it is laid out comes
from the `kind: excel` source of ~/.parley/myorganizer/profile.yaml — nothing about a
particular user or employer is written here.

Usage:
    python objectives.py [--profile PATH] [--year YYYY] [--file PATH.xlsx]

Output: {year, file, categories:[{name, weight, achieved, objectives:[...]}],
         reviews:{"<sheet>": [...]}}

Layout the reader expects (each piece configurable in the profile source):
  sheet         the sheet with the signed objectives                 (default: first)
  review_sheets regex of the periodic review sheets                  (default: none)
  columns       {text, weight, extra} column letters                 (default: A, B, C)
  start_marker  text that opens the table, "{year}" allowed          (default: top)
  start_weight  what the weight column holds on that header row — tells the table's
                header from a loose title with the same words             (optional)
  end_marker    text that closes the table                           (default: bottom)
A row with a weight is a CATEGORY when its text is a short label, an OBJECTIVE when it
is a sentence; weights are fractions (0..1). `extra` holds "achieved" on a category and
the review comment on an objective.

Needs openpyxl and PyYAML (`pip install parley[organizer]`).
"""

import argparse
import glob
import json
import os
import re
import shutil
import sys
import tempfile
from datetime import date

DEFAULT_PROFILE = os.path.join(os.path.expanduser("~"), ".parley", "myorganizer",
                               "profile.yaml")


def load_source(profile_path):
    """The first `kind: excel` source of the profile, or {} when there is none."""
    import yaml
    if not os.path.exists(profile_path):
        return {}
    with open(profile_path, encoding="utf-8") as f:
        profile = yaml.safe_load(f) or {}
    for source in profile.get("sources") or []:
        if source.get("kind") == "excel":
            return source
    return {}


def find_workbook(source, year):
    """Resolve `path` (a file or a folder, "{year}" allowed) to one .xlsx."""
    raw = source.get("path")
    if not raw:
        raise SystemExit("The profile has no excel source with a `path`; pass --file.")
    path = os.path.expanduser(raw.replace("{year}", str(year)))
    if os.path.isfile(path):
        return path
    if not os.path.isdir(path):
        raise SystemExit(f"Objectives folder not found: {path}")
    found = sorted(p for p in glob.glob(os.path.join(path, "*.xlsx"))
                   if not os.path.basename(p).startswith("~$"))
    if not found:
        raise SystemExit(f"No .xlsx in {path}")
    return found[0]


def load_workbook(path):
    """Open the workbook. If OneDrive or Excel holds a lock, read a temp copy."""
    import openpyxl
    try:
        return openpyxl.load_workbook(path, data_only=True)
    except PermissionError:
        copy = os.path.join(tempfile.gettempdir(),
                            "objectives-" + os.path.basename(path))
        shutil.copyfile(path, copy)
        return openpyxl.load_workbook(copy, data_only=True)


def _rows(ws, cols):
    """[(row, text, weight, extra)] for the rows that have text."""
    out = []
    for r in range(1, ws.max_row + 1):
        text = ws[f"{cols['text']}{r}"].value
        if text is None or str(text).strip() == "":
            continue
        out.append((r, str(text).strip(), ws[f"{cols['weight']}{r}"].value,
                    ws[f"{cols['extra']}{r}"].value))
    return out


def parse_sheet(ws, cols, start_marker=None, end_marker=None, start_weight=None):
    """Categories and their objectives from one sheet (signed or a review)."""
    rows = _rows(ws, cols)
    start, end = 0, ws.max_row + 1
    if start_marker:
        hits = [r for r, text, weight, _ in rows
                if text.upper().startswith(start_marker.upper())
                and (start_weight is None
                     or str(weight).strip().upper() == str(start_weight).upper())]
        if not hits:
            return []
        start = hits[0]
    if end_marker:
        hits = [r for r, text, _, _ in rows
                if r > start and text.upper().startswith(end_marker.upper())]
        end = hits[0] if hits else end

    categories = []
    for r, text, weight, extra in rows:
        if not (start < r < end):
            continue
        # Weights are fractions (0.25 / 0.5 / 1). Anything else is not a useful row.
        if not isinstance(weight, (int, float)) or isinstance(weight, bool):
            continue
        if not 0 <= weight <= 1:
            continue
        # Category: a short label ("Delivery"). Objective: a sentence.
        if len(text) < 60 and not text.endswith("."):
            categories.append({"name": text, "weight": weight, "achieved": extra,
                               "objectives": []})
        elif categories:
            categories[-1]["objectives"].append(
                {"text": text, "relative_weight": weight, "comment": extra})
    return categories


def read(source, year, file=None):
    path = file or find_workbook(source, year)
    wb = load_workbook(path)
    cols = {"text": "A", "weight": "B", "extra": "C", **(source.get("columns") or {})}
    start = (source.get("start_marker") or "").replace("{year}", str(year)) or None
    end = source.get("end_marker") or None
    header = source.get("start_weight")

    data = {"year": year, "file": str(path), "categories": [], "reviews": {}}
    sheet = source.get("sheet")
    signed = wb[sheet] if sheet in wb.sheetnames else wb.worksheets[0]
    data["categories"] = parse_sheet(signed, cols, start, end, header)

    pattern = source.get("review_sheets")
    if pattern:
        for name in wb.sheetnames:
            if re.fullmatch(pattern, name):
                data["reviews"][name] = parse_sheet(wb[name], cols, start, end, header)
    return data


def main():
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--profile", default=DEFAULT_PROFILE)
    ap.add_argument("--year", type=int, default=date.today().year)
    ap.add_argument("--file", help="A specific .xlsx — e.g. a copy, when a sandbox "
                                   "or OneDrive blocks the original.")
    args = ap.parse_args()
    try:
        source = load_source(args.profile)
        data = read(source, args.year, args.file)
    except ImportError as err:
        raise SystemExit(f"{err.name} is missing: pip install parley[organizer]")
    # A Windows console is cp1252: accents arrive garbled and anything else raises.
    # The reader is a model that expects UTF-8, so say so on the stream itself.
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    json.dump(data, sys.stdout, ensure_ascii=False, indent=2, default=str)
    sys.stdout.write("\n")


if __name__ == "__main__":
    main()
